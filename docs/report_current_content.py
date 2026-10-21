"""Academic report for the QLCT Android application.

Screenshots are actual emulator captures from a separate demo installation.
Member assignments are proposals until confirmed by the team.
"""
from pathlib import Path
import json, subprocess
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
S='docs/screenshots/current/'
D='docs/diagrams/current/'

def front(r):
    from build_final_report import MEMBERS,field
    metadata_path=ROOT/'docs/report_metadata.json'
    meta=json.loads(metadata_path.read_text(encoding='utf-8')) if metadata_path.exists() else {}
    def cover(text,size=15):
        p=r.doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=p.add_run(text); run.bold=True; run.font.size=Pt(size)
    cover(meta.get('school','[BỔ SUNG TÊN TRƯỜNG]'))
    cover(meta.get('faculty','[BỔ SUNG KHOA / BỘ MÔN]'))
    r.p('')
    cover('BÁO CÁO ĐỒ ÁN\nMÔN LẬP TRÌNH DI ĐỘNG',21)
    cover('ỨNG DỤNG QUẢN LÝ\nCHI TIÊU CÁ NHÂN',24)
    cover('QLCT • Android • Kotlin • Compose • Room SQLite',12)
    r.p('Giảng viên hướng dẫn: '+meta.get('lecturer','[Bổ sung họ tên giảng viên]'))
    r.p('Lớp: '+meta.get('class','[Bổ sung tên lớp]'))
    r.table('Danh sách nhóm',['STT','Họ và tên','Mã sinh viên'],[(i,n,s) for i,(n,s) in enumerate(MEMBERS,1)],[1.2,9.8,5])
    cover('Năm học: '+meta.get('academic_year','[Bổ sung năm học]'),12)
    cover('Tháng 10 năm 2026',12)
    r.chapter('TÓM TẮT ĐỒ ÁN')
    r.p('Đề tài xây dựng ứng dụng quản lý chi tiêu cá nhân trên Android nhằm hỗ trợ ghi nhận thu nhập, theo dõi các khoản chi và kiểm soát ngân sách tháng. Hệ thống gồm tài khoản, quản lý giao dịch, tra cứu, thống kê và ngân sách; giao diện hỗ trợ chế độ sáng và tối.')
    r.p('Mỗi tài khoản quản lý dữ liệu thu/chi và hạn mức riêng trên cùng thiết bị. Dữ liệu được lưu bằng SQLite thông qua Room. Giao diện sử dụng Jetpack Compose và Material 3; mã nguồn được tổ chức theo kiến trúc MVVM. Hai loại biểu đồ được vẽ bằng Canvas để biểu diễn phân bổ và xu hướng chi tiêu.')
    r.p('Báo cáo trình bày bài toán, yêu cầu chức năng, thiết kế tổng thể, cơ sở dữ liệu, quá trình xử lý nghiệp vụ và giao diện sử dụng. Các hình minh họa gồm sơ đồ thiết kế và ảnh chụp ứng dụng với dữ liệu minh họa trên máy ảo Android.')
    r.chapter('MỤC LỤC')
    field(r.doc.add_paragraph(),' TOC \\o "1-3" \\h \\z \\u ','Mở Word, nhấn Ctrl+A rồi F9 để cập nhật mục lục.')
    r.chapter('DANH MỤC HÌNH'); r.figure_index=r.doc.add_paragraph()
    r.chapter('DANH MỤC BẢNG'); r.table_index=r.doc.add_paragraph()
    r.chapter('DANH MỤC TỪ VIẾT TẮT')
    r.table('Các thuật ngữ sử dụng',['Thuật ngữ','Giải thích'],[
        ('UI','Giao diện người dùng.'),('MVVM','Model – View – ViewModel: tách giao diện, trạng thái và xử lý dữ liệu.'),
        ('DAO','Data Access Object: interface truy cập dữ liệu qua Room.'),('CRUD','Create, Read, Update, Delete: thêm, đọc, sửa, xóa.'),
        ('SQLite / Room','SQLite là cơ sở dữ liệu trên thiết bị; Room là lớp truy cập dữ liệu phía trên SQLite.'),
        ('PK / FK','Khóa chính / khóa ngoại trong cơ sở dữ liệu quan hệ.'),('ERD','Sơ đồ quan hệ thực thể.'),
        ('Flow / StateFlow','Luồng bất đồng bộ / luồng trạng thái mới nhất cho giao diện.'),('KSP','Công cụ sinh mã, được dùng cho Room.'),
        ('AVD / APK','Thiết bị Android ảo / gói cài đặt ứng dụng Android.'),('PBKDF2 / salt','Thuật toán dẫn xuất khóa mật khẩu / dữ liệu ngẫu nhiên riêng của mật khẩu.'),
        ('VND','Đồng Việt Nam; ứng dụng nhập và lưu số nguyên.')],[3.2,12.8])

def analysis(r):
    r.chapter('CHƯƠNG 1. GIỚI THIỆU ĐỀ TÀI')
    r.h('1.1. Lý do chọn đề tài')
    r.p('Việc ghi nhớ những khoản tiền nhỏ hằng ngày thường không chính xác. Người dùng có thể nhớ số tiền lớn như tiền phòng hoặc lương nhưng bỏ sót chi phí ăn uống, đi lại và mua sắm. Khi cần xem lại một tháng, dữ liệu phân tán trong ghi chú và trí nhớ khiến việc xác định nguyên nhân chi tiêu tăng trở nên khó khăn. Một ứng dụng ghi thu/chi ngay trên điện thoại giúp tập trung dữ liệu và hỗ trợ nhìn lại thói quen tài chính cá nhân.')
    r.p('Đề tài phù hợp môn Lập trình di động vì kết hợp xây dựng giao diện, xử lý sự kiện, quản lý trạng thái, lưu trữ quan hệ và biểu diễn dữ liệu. Phạm vi cục bộ giúp nhóm tập trung vào nghiệp vụ và SQLite, đồng thời trình diễn được mà không phụ thuộc kết nối Internet. Các chức năng ngân sách và thống kê làm rõ ý nghĩa của dữ liệu sau khi nhập, thay vì chỉ dừng ở một danh sách CRUD.')
    r.h('1.2. Mục tiêu và đối tượng sử dụng')
    r.p('Mục tiêu là xây dựng ứng dụng Android có luồng thao tác dễ hiểu, lưu dữ liệu bền vững và cho phép người dùng trả lời ba câu hỏi: đã thu/chi bao nhiêu, chi nhiều ở danh mục nào và mức chi tháng có vượt hạn mức hay chưa. Đối tượng sử dụng là sinh viên hoặc cá nhân muốn ghi nhận giao dịch thủ công trên một thiết bị.')
    for text in ['Ghi nhận và quản lý khoản thu/chi với số tiền, danh mục, ngày và ghi chú.', 'Tách dữ liệu theo tài khoản cục bộ; cho phép đăng ký, đăng nhập và đăng xuất.', 'Tra cứu bằng từ khóa tiếng Việt có dấu hoặc không dấu, kết hợp tháng và loại giao dịch.', 'Trực quan hóa phân bổ chi theo danh mục và mức chi trong sáu tháng.', 'Đặt hạn mức chi tiêu riêng cho từng tháng và từng tài khoản.', 'Áp dụng Room SQLite, MVVM, Flow và kiểm thử để làm rõ cách tổ chức một ứng dụng hoàn chỉnh.']: r.bullet(text)
    r.h('1.3. Phạm vi và giới hạn')
    r.p('Ứng dụng ghi nhận giao dịch do người dùng nhập. Danh mục được khởi tạo sẵn gồm Ăn uống, Di chuyển, Học tập, Mua sắm, Sinh hoạt, Lương / Trợ cấp và Khác. Bản hiện tại chưa có giao diện tự thêm/sửa/xóa danh mục, nhiều ví tiền, giao dịch định kỳ, xuất Excel/PDF, quét hóa đơn, liên kết ngân hàng hoặc thông báo hệ thống khi vượt ngân sách.')
    r.p('Chênh lệch thu − chi là kết quả của các giao dịch trong kỳ được chọn, không tự động đại diện cho toàn bộ tài sản hoặc số dư ngân hàng của người dùng. Đăng nhập là xác thực cục bộ; dữ liệu không được đồng bộ qua các thiết bị. Gỡ ứng dụng hoặc xóa dữ liệu ứng dụng có thể làm mất dữ liệu vì chưa có chức năng sao lưu/khôi phục do người dùng điều khiển.')
    r.h('1.4. Phương pháp thực hiện')
    r.steps(['Phân tích nghiệp vụ ghi thu/chi, xác định dữ liệu và các tình huống lỗi.', 'Thiết kế bốn bảng, quan hệ khóa ngoại và truy vấn theo tài khoản.', 'Tổ chức mã theo MVVM; xây giao diện Compose và biểu mẫu Material 3.', 'Triển khai CRUD, bộ lọc, tìm kiếm, công thức tổng hợp và Canvas.', 'Bổ sung ngân sách và kiểm tra dữ liệu; chạy kiểm thử trên JVM, SQLite máy tính và Android.', 'Chụp ảnh thật, đối chiếu báo cáo với mã nguồn và chuẩn bị kịch bản bảo vệ.'])
    r.image(D+'31_current_usecase.png','Sơ đồ Use Case tổng quát của ứng dụng')

    r.chapter('CHƯƠNG 2. PHÂN TÍCH YÊU CẦU')
    r.h('2.1. Tác nhân và các nhóm chức năng')
    r.p('Tác nhân chính là người sử dụng trên thiết bị Android. Người chưa xác thực chỉ thao tác với đăng ký và đăng nhập. Sau khi đăng nhập, ứng dụng tạo ExpenseRepository gắn với userId của tài khoản để đọc và thay đổi dữ liệu tương ứng. Bản hiện tại không có tác nhân quản trị hoặc quyền quản trị riêng.')
    r.table('Yêu cầu chức năng và thành phần triển khai',['Mã','Yêu cầu','Mã nguồn liên quan'],[
        ('FR01','Đăng ký, kiểm tra tên và mật khẩu, chống trùng tên.','AuthScreen, AuthViewModel, AuthRepository, AccountDao'),
        ('FR02','Đăng nhập, xác minh mật khẩu, đăng xuất.','AuthRepository.login, AuthViewModel.logout'),
        ('FR03','Thêm khoản thu hoặc chi vào tài khoản hiện tại.','TransactionEditor, ExpenseRepository.save, ExpenseDao.insert'),
        ('FR04','Sửa giao dịch thuộc tài khoản; giữ ID.','TransactionEditor, updateOwned'),
        ('FR05','Xác nhận rồi xóa giao dịch thuộc tài khoản.','TransactionTile, ExpenseScreen, deleteOwned'),
        ('FR06','Danh sách theo ngày giảm dần và ID giảm dần.','ExpenseDao.observeTransactions'),
        ('FR07','Lọc tháng/toàn bộ và ALL/INCOME/EXPENSE.','ExpenseViewModel, summarize'),
        ('FR08','Tìm ghi chú hoặc danh mục không phân biệt dấu.','searchTransactions'),
        ('FR09','Tổng thu, tổng chi, chênh lệch; biểu đồ vòng và cột.','ExpenseSummary, DashboardHeader, ExpenseCharts'),
        ('FR10','Đặt/sửa/xóa ngân sách theo tài khoản và tháng.','MonthlyBudget, BudgetEditor, saveBudget, deleteBudget'),
        ('FR11','Cảnh báo 80%, chạm và vượt hạn mức.','BudgetStatus, MonthlyBudgetCard'),
        ('FR12','Lưu và đọc dữ liệu khi đóng, mở ứng dụng.','ExpenseDatabase, Room và SQLite')],[1.3,6.7,8])
    r.h('2.2. Yêu cầu phi chức năng')
    r.table('Yêu cầu chất lượng',['Nhóm','Yêu cầu và cách đáp ứng'],[
        ('Dễ sử dụng','Nhãn tiếng Việt, tông tím thống nhất, thanh dưới chia đều; nút + nổi bật ở chính giữa.'),
        ('Chính xác','Tiền lưu Long; kiểm tra giá trị dương, giới hạn tiền và ngày trước khi ghi.'),
        ('Phản hồi thao tác','Có busy, actionError, loadError; giữ biểu mẫu khi lưu lỗi; chặn gửi lặp khi đang xử lý.'),
        ('Lưu trữ','Room trên SQLite qlct.db; định nghĩa khóa chính, chỉ mục và khóa ngoại.'),
        ('Phân tách tài khoản','Truy vấn có điều kiện userId; sửa/xóa dùng cả ID giao dịch và chủ sở hữu.'),
        ('Mật khẩu','Lưu salt và hash PBKDF2 thay cho mật khẩu rõ; đây không phải mã hóa toàn bộ database.'),
        ('Khả năng học và bảo trì','Tách View, ViewModel, Repository, DAO; hàm tổng hợp và validator có unit test.'),
        ('Tương thích','minSdk 24, targetSdk/compileSdk 37 theo cấu hình dự án; bằng chứng chạy trong báo cáo trên API 35.')],[3.5,12.5])
    r.h('2.3. Quy tắc nghiệp vụ')
    r.table('Quy tắc kiểm tra và tính toán',['Nội dung','Quy tắc hiện tại'],[
        ('Tên đăng nhập','Trim, chuyển chữ thường để tra cứu; 3–32 ký tự chữ không dấu, số hoặc dấu _; username UNIQUE.'),
        ('Mật khẩu','8–128 ký tự, không chỉ toàn khoảng trắng; đăng ký phải khớp xác nhận.'),
        ('Số tiền','Số nguyên từ 1 đến 999.999.999.999 VND; type xác định thu hoặc chi, không nhập chi bằng số âm.'),
        ('Ngày / tháng','Ngày ISO yyyy-MM-dd; tháng ngân sách yyyy-MM; giao diện hiển thị ngày dd/MM/yyyy.'),
        ('Danh mục','Phải chọn danh mục tồn tại; cơ sở dữ liệu có FK category_id.'),
        ('Tổng kỳ','Chỉ phụ thuộc tháng/toàn bộ, không phụ thuộc lọc Thu/Chi hoặc từ khóa.'),
        ('Biểu đồ vòng','Chỉ EXPENSE của kỳ; gom theo danh mục; sắp số tiền giảm dần.'),
        ('Biểu đồ cột','Sáu tháng liên tiếp kết thúc ở tháng chọn; chỉ EXPENSE; không phụ thuộc từ khóa.'),
        ('Ngân sách','Một hạn mức cho mỗi cặp tài khoản–tháng; tính chi đúng tháng kể cả đang xem toàn bộ.'),
        ('Xóa dữ liệu','Xóa giao dịch cần xác nhận; xóa ngân sách chỉ xóa hạn mức, không xóa giao dịch.')],[3.5,12.5])
    r.h('2.4. Đặc tả các luồng tiêu biểu')
    for title,pre,steps,post in [
        ('UC01 — Đăng ký và đăng nhập','Ứng dụng đã đọc hoặc khởi tạo dữ liệu tài khoản.', ['Người dùng chọn Đăng ký, nhập tên, mật khẩu và xác nhận.', 'ViewModel kiểm tra định dạng; Repository kiểm tra trùng tên và tạo salt/hash.', 'Lưu tài khoản; thành công mở ExpenseScreen theo userId mới. Nếu đã có tài khoản, dùng Đăng nhập và xác minh hash.', 'Sai thông tin hiển thị lỗi; người dùng có thể sửa rồi gửi lại.'], 'Tài khoản mới bắt đầu với giao dịch rỗng; riêng vu98 có dữ liệu demo một lần.'),
        ('UC02 — Thêm hoặc sửa giao dịch','Người dùng đã đăng nhập, danh mục tải thành công.', ['Nhấn + để thêm hoặc chọn dòng để sửa.', 'Chọn Thu/Chi, nhập tiền, danh mục, ngày và ghi chú.', 'Nhấn Lưu; lỗi nhập giữ hộp thoại và thông báo tại ô tiền.', 'ViewModel gọi Repository; kiểm tra lại, gắn chủ sở hữu rồi insert/updateOwned.'], 'Thành công: Flow cập nhật dữ liệu và hộp thoại đóng. Lỗi lưu: giữ dữ liệu để thử lại.'),
        ('UC03 — Tra cứu giao dịch','Có thể có hoặc không có dữ liệu giao dịch.', ['Mở Giao dịch, chọn tháng hoặc Toàn bộ.', 'Chọn Tất cả/Thu/Chi và nhập từ khóa.', 'Ứng dụng bỏ dấu, chuẩn hóa chữ thường và tìm trong ghi chú/tên danh mục.', 'Nếu không có kết quả, hiển thị hướng dẫn thay đổi từ khóa hoặc bộ lọc.'], 'Danh sách thay đổi nhưng các tổng và biểu đồ vẫn theo quy tắc thống kê của kỳ.'),
        ('UC04 — Quản lý ngân sách','Người dùng đã đăng nhập và đang chọn một tháng.', ['Mở Ngân sách, chọn Đặt hoặc Sửa ngân sách.', 'Nhập số nguyên dương và lưu; @Upsert thêm mới hoặc cập nhật theo khóa ghép.', 'Quan sát số đã chi, còn lại/vượt và tiến trình.', 'Nếu xóa, xác nhận xóa hạn mức; các giao dịch không bị xóa.'], 'Ngân sách được lưu riêng theo tài khoản và tháng; cảnh báo phản ánh tổng chi tháng.')]:
        r.h(title,3); r.p('Điều kiện trước: '+pre); r.steps(steps); r.p('Kết quả: '+post)

def design(r):
    r.chapter('CHƯƠNG 3. CÔNG NGHỆ VÀ THIẾT KẾ HỆ THỐNG')
    r.h('3.1. Công nghệ sử dụng')
    r.table('Thành phần công nghệ trong project',['Thành phần','Vai trò'],[
        ('Kotlin','Ngôn ngữ viết ứng dụng; data class, coroutine, xử lý null và các hàm tổng hợp.'),
        ('Jetpack Compose / Material 3','Giao diện bằng @Composable: Scaffold, LazyColumn, AlertDialog, DatePicker, TextField, Button.'),
        ('Canvas của Compose','Tự vẽ biểu đồ vòng, cột và biểu tượng nét; không thêm thư viện chart bên ngoài.'),
        ('ViewModel / Lifecycle 2.9.4','Giữ trạng thái màn hình; collectAsStateWithLifecycle thu nhận state theo lifecycle.'),
        ('Room 2.8.4','Entity, DAO, kiểm tra truy vấn và sinh mã SQLite; Flow phản ánh thay đổi dữ liệu.'),
        ('Coroutines / Flow','Hàm suspend cho truy cập dữ liệu; combine, stateIn và luồng trạng thái bất đồng bộ.'),
        ('KSP 2.3.6','Xử lý annotation để sinh phần triển khai Room.'),
        ('Android Gradle Plugin 9.4.1','Build ứng dụng; Kotlin Compose plugin 2.2.10; Compose BOM 2025.08.00.'),
        ('JUnit / AndroidJUnitRunner','Unit test trên JVM và kiểm thử tích hợp trên Android.')],[4.8,11.2])
    r.p('SQLite đảm nhiệm lưu trữ dữ liệu quan hệ trên thiết bị. Room ánh xạ bảng với đối tượng Kotlin, khai báo truy vấn qua DAO và hỗ trợ quan sát dữ liệu bằng Flow. Tài khoản, danh mục, giao dịch và ngân sách được lưu trong cùng cơ sở dữ liệu qlct.db. [1]')
    r.p('ViewModel cung cấp trạng thái cho giao diện và giữ dữ liệu trạng thái qua thay đổi cấu hình. Trong QLCT, các lựa chọn tháng, monthOnly, type và query được đặt trong SavedStateHandle; việc đăng nhập không được mô tả như một phiên đăng nhập vĩnh viễn vì trạng thái xác thực hiện không lưu thành token bền vững. [2][3]')
    r.h('3.2. Kiến trúc và luồng dữ liệu')
    r.image(D+'21_current_architecture.png','Các tầng View, ViewModel, Repository, DAO và Room SQLite')
    r.p('MainActivity mở database singleton và lắp ghép các dependency bằng tay. AuthRepository nhận AccountDao, DatabaseInitializer và DemoDataSeeder; AuthViewModel điều khiển trang xác thực. Khi có userId, MainActivity tạo ExpenseRepository theo tài khoản và ExpenseViewModel tương ứng. key(authState.userId) giúp phân tách cây giao diện khi đổi tài khoản.')
    r.p('ExpenseScreen quan sát một ExpenseUiState bất biến. Các sự kiện như save, delete và saveBudget chuyển về ViewModel thay vì ghi database trực tiếp trong Composable. ViewModel kết hợp các luồng danh mục, giao dịch, ngân sách và lựa chọn lọc để tính state. Khi Room phát dữ liệu mới, Compose nhận state và hiển thị lại phần có thay đổi.')
    r.table('Trách nhiệm theo package',['Thư mục dưới com/example/qlct','Trách nhiệm'],[
        ('view/auth','Trang đăng nhập/đăng ký.'),('view/expense','Màn hình chính, biểu mẫu giao dịch và ngân sách.'),
        ('view/components','Header, thanh điều hướng, thẻ, dòng giao dịch, biểu đồ và định dạng tiền.'),('view/theme','Màu sắc và theme Compose.'),
        ('viewmodel và viewmodel/state','Sự kiện, state, bộ lọc, busy/error và phiên xác thực trong bộ nhớ.'),
        ('model/entity và model/dao','Bảng dữ liệu và interface truy vấn Room.'),('model/database','Mở database, khởi tạo danh mục/tài khoản và demo.'),
        ('model/repository','Điểm truy cập nghiệp vụ; gắn userId và kiểm tra trước khi ghi.'),('model/validation và model/summary','Kiểm tra đầu vào, tính tổng hợp, tìm kiếm và trạng thái ngân sách.')],[6,10])
    r.image(D+'28_current_navigation.png','Điều hướng bốn trang với nút thêm giao dịch ở chính giữa')
    r.p('Điều hướng hiện dùng biến tab được rememberSaveable, không dùng NavHost hoặc một back stack riêng. Các giá trị 0, 1, 2, 3 lần lượt là Tổng quan, Giao dịch, Thống kê, Ngân sách. Nút + mở AlertDialog TransactionEditor. Các nút dùng vùng bố trí có weight bằng nhau để dấu + nằm giữa chiều rộng thanh dưới.')
    r.h('3.3. Thiết kế cơ sở dữ liệu')
    r.image(D+'27_current_erd.png','Sơ đồ quan hệ bốn bảng trong cơ sở dữ liệu qlct.db')
    r.p('users là bảng tài khoản cục bộ. Mỗi tài khoản có ID tự tăng, username duy nhất, salt và passwordHash. demoSeeded đánh dấu đã tạo các giao dịch minh họa cho vu98. categories gồm ID và tên duy nhất; danh mục dùng chung giữa các tài khoản. transactions tham chiếu danh mục và tài khoản; budgets tham chiếu tài khoản và có khóa chính ghép theo tháng.')
    schemas=[('users',[
        ('id','INTEGER / Int','PK, tự tăng','Định danh tài khoản.'),('username','TEXT','NOT NULL, UNIQUE','Tên chuẩn hóa dùng đăng nhập.'),('salt','TEXT','NOT NULL','Salt dạng hex.'),('passwordHash','TEXT','NOT NULL','Hash dẫn xuất, không phải mật khẩu rõ.'),('demoSeeded','INTEGER / Boolean','NOT NULL, mặc định 0','Đánh dấu seed demo đã hoàn tất.')]),
        ('categories', [('id','INTEGER / Long','PK','ID danh mục cố định.'),('name','TEXT','NOT NULL, UNIQUE','Tên danh mục tiếng Việt.')]),
        ('transactions', [('id','INTEGER / Long','PK, tự tăng','ID giao dịch.'),('type','TEXT','NOT NULL','INCOME hoặc EXPENSE, kiểm tra ở Repository.'),('amount','INTEGER / Long','NOT NULL','Số tiền VND nguyên dương.'),('category_id','INTEGER / Long','NOT NULL, FK','Tham chiếu categories.id.'),('date','TEXT','NOT NULL','Ngày ISO yyyy-MM-dd.'),('note','TEXT','NOT NULL, mặc định rỗng','Ghi chú tùy chọn.'),('user_id','INTEGER / Int?','Có thể NULL, FK','Tài khoản sở hữu; cột cho phép NULL theo định nghĩa lưu trữ.')]),
        ('budgets', [('user_id','INTEGER / Int','PK ghép, NOT NULL, FK','Tài khoản sở hữu.'),('month','TEXT','PK ghép, NOT NULL','Tháng yyyy-MM.'),('amount','INTEGER / Long','NOT NULL','Hạn mức; validator yêu cầu lớn hơn 0.')])]
    for name,rows in schemas:
        r.h('Bảng '+name,3)
        r.table('Từ điển dữ liệu '+name,['Cột','Kiểu dữ liệu','Ràng buộc','Ý nghĩa'],rows,[3,3.2,4,5.8])
    r.p('Các khóa ngoại sử dụng NO ACTION khi cập nhật hoặc xóa bản ghi cha. Không có cơ chế cascade xóa tài khoản/danh mục trong bản hiện tại. transactions có chỉ mục trên date, category_id và user_id; username và tên danh mục có chỉ mục UNIQUE. budgets dùng chỉ mục khóa chính ghép. Các chỉ mục hỗ trợ truy vấn nhưng không có nghĩa mọi câu JOIN hoặc sắp xếp đều không cần xử lý bổ sung.')
    r.h('3.4. Tổ chức lưu trữ và truy cập dữ liệu')
    r.p('ExpenseDatabase khai báo bốn thực thể User, Category, TransactionEntity và MonthlyBudget. AccountDao thực hiện thao tác tài khoản; ExpenseDao xử lý danh mục, giao dịch và ngân sách. Repository kiểm tra dữ liệu nghiệp vụ trước khi chuyển yêu cầu xuống DAO. Các bảng sử dụng khóa chính, khóa ngoại và chỉ mục để tổ chức dữ liệu và hỗ trợ truy vấn.')
    r.p('Danh mục được chuẩn bị sẵn để người dùng chọn khi nhập giao dịch. Mỗi giao dịch lưu loại thu/chi, số tiền, danh mục, ngày, ghi chú và tài khoản sở hữu. Ngân sách được lưu theo tài khoản và tháng. Khi dữ liệu thay đổi, Room phát kết quả qua Flow để cập nhật danh sách, thẻ tổng và biểu đồ.')
    r.h('3.5. Thiết kế tổng thể hệ thống')
    r.image(D+'26_current_system.png','Mô hình tổng thể hệ thống quản lý chi tiêu cá nhân')
    r.p('Người dùng tương tác qua các màn hình xác thực, tổng quan, giao dịch, thống kê và ngân sách. Tầng xử lý kiểm tra dữ liệu nhập, giới hạn dữ liệu theo tài khoản, tổng hợp số liệu và xác định cảnh báo. Tầng lưu trữ cung cấp dữ liệu qua Room SQLite. Chu trình thao tác gồm nhập yêu cầu, kiểm tra, đọc hoặc ghi dữ liệu, tính toán và hiển thị kết quả.')
    r.h('3.6. Xác thực và phân tách tài khoản')
    r.image(D+'22_current_auth.png','Luồng xác thực cục bộ và seed demo sau đăng nhập')
    r.p('Khi đăng ký, AuthRepository sinh salt 32 byte bằng SecureRandom rồi dẫn xuất hash 256 bit với PBKDF2WithHmacSHA1, 600.000 vòng. Mật khẩu không được lưu dạng rõ. Khi đăng nhập, ứng dụng dẫn xuất lại với salt đã lưu và so sánh bằng MessageDigest.isEqual. Công việc dẫn xuất chạy ở Dispatchers.Default để tránh thực hiện trực tiếp trên luồng giao diện.')
    r.p('ExpenseRepository nhận userId từ trạng thái xác thực, gắn userId khi thêm và gọi updateOwned/deleteOwned khi sửa/xóa. Các truy vấn có cả điều kiện ID và chủ sở hữu. DAO vẫn có một số hàm thao tác tổng quát phục vụ kiểm thử/lịch sử; luồng nghiệp vụ hiện tại phải đi qua Repository để duy trì kiểm tra và giới hạn tài khoản. SQLite hiện chưa được mã hóa toàn bộ; không nên diễn giải việc hash mật khẩu thành mã hóa dữ liệu thu/chi.')

def implementation(r):
    r.chapter('CHƯƠNG 4. CÀI ĐẶT CÁC CHỨC NĂNG')
    r.h('4.1. Thêm, sửa và xóa giao dịch')
    r.image(D+'23_current_save.png','Luồng lưu giao dịch từ hộp thoại tới SQLite và cập nhật giao diện')
    r.p('TransactionEditor dùng trạng thái Compose cho type, amount, categoryId, date và note. Khi nhấn Lưu, amountChecked được bật để hiển thị lỗi tại ô tiền. AmountValidator kiểm tra chuỗi rỗng, số nguyên, giá trị dương và giới hạn tối đa. Bộ chọn ngày dùng DatePicker; giá trị ngày được chuyển đổi theo UTC để bảo toàn ngày được chọn khi trao đổi với thành phần chọn ngày.')
    r.p('ViewModel.mutate bỏ qua thao tác mới nếu busy đang true, bật busy trước khi chạy coroutine và đặt lại trong finally. Khi thành công, completedActions tăng; ExpenseScreen dùng thay đổi này để đóng biểu mẫu. Khi có lỗi, actionError được đưa vào UiState và dữ liệu nhập được giữ để thử lại. CancellationException được truyền tiếp thay vì biến thành lỗi nghiệp vụ.')
    r.p('Repository kiểm tra lại type, số tiền, categoryId và định dạng ngày. Ghi chú được trim, userId luôn lấy từ Repository. Với ID bằng 0, DAO insert và sinh ID; với ID đã có, updateOwned chỉ thay đổi bản ghi của tài khoản hiện tại. Nếu số dòng cập nhật không bằng 1, thao tác được xem là thất bại. Xóa cũng kiểm tra số dòng và cần xác nhận ở giao diện.')
    r.code('TransactionEditor → ExpenseViewModel.save(transaction)\n→ mutate { repository.save(transaction) }\n→ validate + transaction.copy(note = note.trim(), userId = userId)\n→ id == 0: dao.insert(...)\n→ id != 0: dao.updateOwned(id, userId, ...)\n→ Room Flow → UiState → Compose')
    r.h('4.2. Truy vấn, lọc và tìm kiếm')
    r.code('SELECT t.id,t.type,t.amount,t.category_id,t.date,t.note,\n       c.name AS categoryName\nFROM transactions t\nJOIN categories c ON c.id = t.category_id\nWHERE t.user_id IS :userId\nORDER BY t.date DESC, t.id DESC')
    r.p('DAO trả Flow<List<TransactionRow>>, trong đó TransactionRow ghép tên danh mục để giao diện không phải tự tra cứu từng dòng. Việc lưu ngày theo ISO giúp so sánh/sắp xếp chuỗi đúng thứ tự thời gian đối với các ngày hợp lệ. Tổng hợp hiện thực hiện trên danh sách được quan sát trong bộ nhớ, phù hợp dữ liệu nhỏ của đồ án; khi dữ liệu lớn nên cân nhắc phân trang và tổng hợp bằng truy vấn SQL.')
    r.image(D+'24_current_search.png','Thứ tự áp dụng kỳ, loại giao dịch và từ khóa')
    r.p('summarize lấy period theo monthOnly. Tổng income và expense tính trên period trước khi lọc loại. visible là period sau bộ lọc ALL/INCOME/EXPENSE. searchTransactions tiếp tục xử lý visible để tạo searchResults. Hàm tìm kiếm dùng Unicode Normalizer dạng NFD, bỏ dấu kết hợp, đổi đ thành d và chuyển chữ thường. Câu tìm kiếm được tách thành các từ; mọi từ đều phải xuất hiện trong chuỗi ghép ghi chú và tên danh mục.')
    r.p('Ví dụ, nhập “an uong” có thể tìm danh mục “Ăn uống”. Nhập “demo an” yêu cầu cả “demo” và “an” xuất hiện, không nhất thiết liền nhau. Không tìm theo số tiền hoặc ngày bằng từ khóa. Nếu chọn Thu nhưng từ khóa chỉ khớp khoản chi, kết quả rỗng là hợp lệ; người dùng cần đổi bộ lọc hoặc xóa từ khóa.')
    r.h('4.3. Tổng quan và công thức thống kê')
    r.table('Công thức nghiệp vụ',['Chỉ số','Công thức / phạm vi'],[
        ('Tổng thu','Σ amount với type = INCOME trong kỳ được chọn.'),('Tổng chi','Σ amount với type = EXPENSE trong kỳ được chọn.'),('Chênh lệch','Tổng thu − Tổng chi; có thể âm hợp lệ.'),
        ('Chi theo danh mục','Gom EXPENSE trong kỳ theo categoryId và cộng amount.'),('Tỷ lệ danh mục','Chi danh mục / tổng chi × 100%; nhãn hiện dùng toInt nên có làm tròn xuống.'),
        ('Chi từng tháng','Σ EXPENSE của mỗi tháng trong sáu tháng kết thúc tại tháng chọn.'),('Còn lại ngân sách','Hạn mức tháng − chi riêng tháng; âm nghĩa là vượt hạn mức.')],[4.3,11.7])
    r.p('Ví dụ bộ dữ liệu ảnh báo cáo ban đầu có sáu khoản chi tổng 1.345.000 VND và chưa có khoản thu, nên chênh lệch là −1.345.000 VND. Sau khi thêm khoản thu minh họa 3.000.000 VND, chênh lệch thành 1.655.000 VND. Đây là số của bản demo riêng tại thời điểm chụp, không phải số cố định của mọi tài khoản hoặc mọi lần mở ứng dụng.')
    r.h('4.4. Vẽ biểu đồ bằng Canvas')
    r.image(D+'30_current_charts.png','Cách chuyển số liệu thành góc vòng và chiều cao cột')
    r.p('CategoryChart vẽ vòng bằng Canvas và drawArc với Stroke. Bắt đầu ở góc −90°, mỗi danh mục có sweep = amount / totalExpense × 360. Sau khi vẽ một phần, start được cộng thêm sweep để phần tiếp theo nối tiếp. Nét vòng dày 25 dp; tâm vòng là nhãn Tổng chi và số tiền rút gọn. Không có chi thì vẽ vòng nền để tránh chia cho 0.')
    r.p('SpendingChart lấy sáu MonthlyExpense. Giá trị max được ép tối thiểu bằng 1 để tính tỷ lệ an toàn. Mỗi tháng có một slot theo chiều rộng Canvas; drawRoundRect vẽ cột có chiều cao tỷ lệ với amount/max. Ba đường ngang dùng drawLine; cột tháng cuối màu tím đậm hơn, các cột trước màu tím nhạt. Tháng không có chi vẫn có nhãn nhưng không có cột dương.')
    r.p('Nhãn phần trăm hiện làm tròn xuống thành số nguyên nên tổng các nhãn có thể nhỏ hơn 100%; các góc vẫn dùng tỷ lệ thực. Long được dùng để lưu và cộng tiền, Double/Float chỉ phục vụ phép chia và tọa độ. Canvas là API đồ họa của Compose, không phải một thư viện biểu đồ ngoài như MPAndroidChart. [4]')
    r.h('4.5. Ngân sách tháng')
    r.image(D+'25_current_budget.png','Các trạng thái ngân sách và nguyên tắc chỉ tính chi đúng tháng')
    r.p('MonthlyBudget có khóa ghép user_id và month. saveBudget kiểm tra có chủ sở hữu, tháng đúng định dạng và amount hợp lệ trước khi @Upsert. deleteBudget chỉ xóa bản ghi hạn mức của tài khoản/tháng đó. Ngân sách có thể được mở từ thẻ ở Tổng quan hoặc từ trang Ngân sách ở thanh dưới.')
    r.p('BudgetStatus nhận limit và spent. spent lấy từ cột cuối của chuỗi thống kê sáu tháng, tức đúng tháng đang chọn. Vì vậy dù người dùng bật Toàn bộ hay lọc Thu, số đã chi dùng cho ngân sách vẫn là EXPENSE riêng tháng. Thanh tiến trình được giới hạn trong 0..1; vượt ngân sách không làm thanh kéo dài quá khung.')
    r.table('Ngưỡng cảnh báo và ví dụ với chi 1.345.000 VND',['Trạng thái','Điều kiện','Ví dụ hạn mức'],[
        ('NONE','Chưa có hạn mức hợp lệ.','Chưa đặt ngân sách.'),('NORMAL','Chi dưới 80% hạn mức.','2.000.000 VND → 67,25%.'),
        ('WARNING','Từ 80% và chưa chạm hạn mức.','1.500.000 VND → khoảng 89,67%.'),('REACHED','Chi bằng hạn mức.','1.345.000 VND.'),('EXCEEDED','Chi lớn hơn hạn mức.','1.000.000 VND → vượt 345.000 VND.')],[3.1,6.5,6.4])

def interface(r):
    r.chapter('CHƯƠNG 5. GIAO DIỆN VÀ HƯỚNG DẪN SỬ DỤNG')
    r.p('Các ảnh dưới đây là ảnh chụp thật của bản demo riêng ngày 21/10/2026. Chúng thể hiện mã nguồn ứng dụng, với dữ liệu vu98 minh họa. Bố cục dùng Material 3 và tông tím; nội dung dưới thanh điều hướng có thể cuộn. Ảnh được đặt theo cặp để dễ so sánh các trạng thái. Các ảnh và nội dung số liệu cần được đọc theo thời điểm trong kịch bản, vì khoản thu minh họa được thêm ở các bước sau.')
    pairs=[
        ('5.1. Xác thực tài khoản','01_login','Đăng nhập bằng tài khoản cục bộ','02_register','Đăng ký tài khoản với xác nhận mật khẩu', 'Trang xác thực có ô tên đăng nhập, mật khẩu và nút Ẩn/Hiện. Đăng ký thêm ô xác nhận mật khẩu. Sau xác thực thành công, người dùng chuyển tới giao diện quản lý theo userId. Tài khoản mới thông thường chưa có giao dịch; vu98 là tài khoản demo riêng.'),
        ('5.2. Tổng quan và lịch sử','03_overview','Tổng quan với sáu khoản chi demo','04_transactions','Trang Giao dịch và bộ lọc', 'Tổng quan hiển thị kỳ, tổng thu/chi và chênh lệch, thẻ ngân sách cùng các phần có thể cuộn bên dưới. Giao dịch có tìm kiếm và bộ lọc loại; nhấn dòng để sửa hoặc biểu tượng thùng rác để yêu cầu xóa. Nút + được đặt ở giữa năm vị trí của thanh dưới.'),
        ('5.3. Tra cứu và phân bổ chi','05_search','Tìm “an uong” khớp danh mục tiếng Việt','06_category_statistics','Biểu đồ vòng và số tiền theo danh mục', 'Tìm kiếm hỗ trợ không dấu. Thống kê biểu diễn tỷ lệ các nhóm chi và bổ sung số tiền chi tiết. Danh mục có chi lớn được xếp trước; màu tương ứng giữa chú giải và phần vòng. Không nhầm tổng nhãn phần trăm làm tròn với tổng góc thực tế.'),
        ('5.4. Thống kê tháng và ngân sách chưa đặt','07_month_statistics','Biểu đồ chi trong sáu tháng và số liệu chi tiết','08_budget_empty','Trang Ngân sách khi chưa có hạn mức', 'Biểu đồ cột luôn có sáu tháng kết thúc ở tháng được chọn. Vì dữ liệu demo mới chỉ có chi trong tháng hiện tại, các tháng trước không có cột. Trang Ngân sách tập trung vào hạn mức của tháng, với nút Đặt ngân sách nếu chưa có.'),
        ('5.5. Thiết lập ngân sách và cảnh báo','09_budget_editor','Hộp nhập hạn mức ngân sách','10_budget_warning','Cảnh báo đã dùng từ 80% ngân sách', 'Nhập 1.500.000 VND trong khi đã chi 1.345.000 VND tạo trạng thái WARNING. Đây là cảnh báo hiển thị trong ứng dụng, không phải thông báo đẩy của hệ điều hành. Nút sửa cho phép thay đổi hạn mức của chính tháng đang xem.'),
        ('5.6. Vượt ngân sách và thêm giao dịch','11_budget_exceeded','Ngân sách 1.000.000 VND bị vượt 345.000 VND','12_add_transaction','Hộp Thêm giao dịch', 'Khi vượt hạn mức, ứng dụng hiển thị số tiền vượt và màu cảnh báo. Hộp giao dịch có ô tiền và ghi chú bo góc; ô danh mục và ngày rộng đều nhau; nút Lưu có nền tím. Người dùng chọn Thu hoặc Chi trước khi lưu giao dịch.'),
        ('5.7. Kiểm tra đầu vào và chọn danh mục','13_amount_validation','Báo lỗi khi nhấn Lưu mà chưa nhập số tiền','14_category_picker','Danh mục chọn bằng menu', 'Ví dụ lỗi đầu vào cho thấy hộp thoại vẫn mở và lỗi xuất hiện tại ô số tiền. Danh mục là dữ liệu có sẵn trong categories; menu chọn không phải chức năng thêm danh mục. Người dùng cần nhập số nguyên dương trước khi lưu.'),
        ('5.8. Chọn ngày và ghi khoản thu','15_date_picker','Bộ chọn ngày Material 3','16_income_form','Biểu mẫu thu nhập minh họa 3.000.000 VND', 'Ngày hiển thị dd/MM/yyyy và lưu theo ISO. Khoản thu vẫn nhập số dương rồi chọn Thu tiền, danh mục Lương / Trợ cấp. Ghi chú DEMO_bao_cao_thu_nhap chỉ dùng để nhận diện dữ liệu trình diễn trong bản riêng.'),
        ('5.9. Kết quả lưu và giao diện sửa','17_saved_income','Khoản thu mới xuất hiện trong lịch sử','18_edit_transaction','Mở giao dịch đã lưu để sửa', 'Sau Lưu, dữ liệu mới được Room phát qua Flow và danh sách cập nhật. Mở một giao dịch đã có sẽ điền giá trị cũ vào hộp Sửa giao dịch. Ảnh mở biểu mẫu sửa không tự chứng minh thao tác cập nhật đã được thực hiện; kết quả CRUD được kiểm tra thêm bằng bộ test.'),
        ('5.10. Xác nhận xóa và tổng quan sau khoản thu','19_delete_confirmation','Hộp xác nhận trước khi xóa giao dịch','20_overview_income','Tổng quan sau khi thêm khoản thu demo', 'Trong kịch bản chụp ảnh, hộp xóa được hủy nên khoản thu còn nguyên. Tổng thu 3.000.000 VND và tổng chi 1.345.000 VND cho chênh lệch 1.655.000 VND. Người dùng có thể trở lại Tổng quan bằng nút đầu tiên của thanh dưới.')]
    for title,a,ac,b,bc,text in pairs:
        r.h(title); r.p(text)
        r.pair_images(S+a+'.png',ac,S+b+'.png',bc)
    r.h('5.11. Nút chuyển giao diện sáng/tối')
    r.p('Nút mặt trăng/mặt trời xuất hiện trên trang xác thực và cạnh Đăng xuất. Bấm nút đổi theme ngay, không đăng xuất. Lựa chọn được lưu bằng SharedPreferences và giữ sau khi đóng/mở ứng dụng. Khi chưa chọn, ứng dụng theo chế độ hệ thống; sau khi chọn, giá trị đã lưu được ưu tiên. Màu nền, chữ và mục điều hướng dùng bảng màu Material 3 tương ứng.')
    r.pair_images(S+'22_dark_overview.png','Tổng quan ở chế độ tối',S+'25_light_toggle.png','Ngân sách sau khi chuyển về chế độ sáng')
    r.h('5.12. Trình tự thao tác khuyến nghị')
    r.steps(['Đăng ký tài khoản của riêng mình hoặc dùng tài khoản demo đã thống nhất cho buổi bảo vệ.', 'Từ Tổng quan, chọn đúng tháng cần nhập và xem.', 'Nhấn +, ghi khoản thu/chi, chọn danh mục và ngày, rồi lưu.', 'Mở Giao dịch để tìm lại, đổi bộ lọc, sửa hoặc xóa theo nhu cầu.', 'Mở Thống kê để xem phân bổ và xu hướng chi theo tháng.', 'Mở Ngân sách, nhập hạn mức và theo dõi các ngưỡng cảnh báo.', 'Đăng xuất khi kết thúc; nếu mở lại sau khi tiến trình bị hủy, đăng nhập để truy cập dữ liệu đã lưu.'])

def testing(r):
    data=json.loads((ROOT/'docs/report_validation_current.json').read_text(encoding='utf-8'))
    r.chapter('CHƯƠNG 6. KIỂM THỬ VÀ ĐÁNH GIÁ')
    r.h('6.1. Mục tiêu và môi trường kiểm thử')
    r.p('Kiểm thử đánh giá tính đúng đắn của dữ liệu nhập, thao tác quản lý giao dịch, phân tách tài khoản, kết quả thống kê và cảnh báo ngân sách. Ứng dụng được chạy trên máy ảo Android 15/API 35; các hàm xử lý được kiểm tra bằng unit test và các thao tác lưu trữ được kiểm tra trên Android.')
    r.h('6.2. Kết quả kiểm tra chức năng')
    r.table('Các tình huống kiểm tra tiêu biểu',['Chức năng','Tình huống kiểm tra','Kết quả'],[
        ('Tài khoản','Kiểm tra tên, mật khẩu, xác nhận; xác minh thông tin đăng nhập.','Đạt trong bộ kiểm thử.'),
        ('Giao dịch','Thêm, đọc, sửa, xóa và lưu dữ liệu.','Đạt trong bộ kiểm thử lưu trữ.'),
        ('Số tiền','Rỗng, sai định dạng, không dương và vượt giới hạn.','Hiển thị lỗi, chưa lưu giao dịch.'),
        ('Tìm kiếm','Tìm ghi chú và danh mục bằng từ khóa không dấu.','Hiển thị kết quả phù hợp.'),
        ('Bộ lọc và thống kê','Kỳ thời gian, loại thu/chi, tổng kỳ và chuỗi sáu tháng.','Kết quả đúng với dữ liệu kiểm thử.'),
        ('Ngân sách','Hạn mức theo tài khoản/tháng và các ngưỡng cảnh báo.','Đạt trong bộ kiểm thử.'),
        ('Giao diện sáng/tối','Đổi theme rồi đóng và mở lại ứng dụng.','Đổi màu và giữ lựa chọn.'),
        ('Hủy thao tác','Mở sửa hoặc xác nhận xóa rồi chọn Hủy.','Không thay đổi giao dịch.')],[3.3,8.5,4.2])
    r.h('6.3. Kiểm thử tự động')
    r.p(f"Bộ kiểm thử JVM ghi nhận {data['unit_tests']} trường hợp và bộ kiểm thử Android ghi nhận {data['android_tests']} trường hợp, đều không có lỗi hoặc thất bại. Quá trình biên dịch ứng dụng thành công. Kết quả hỗ trợ đánh giá các quy tắc nghiệp vụ và lưu trữ trong phạm vi được kiểm tra.")
    r.table('Nhóm nội dung kiểm thử',['Nhóm','Phạm vi'],[
        ('Kiểm tra đầu vào','Tên đăng nhập, mật khẩu, số tiền và xác nhận thông tin.'),
        ('Tổng hợp và tìm kiếm','Tổng thu/chi theo kỳ, phân loại, từ khóa tiếng Việt và ngân sách.'),
        ('Trạng thái giao diện','Đang xử lý, thông báo lỗi, hoàn tất và thử lại.'),
        ('Lưu trữ và tài khoản','Đọc/ghi, tính duy nhất, quan hệ bảng và dữ liệu theo chủ sở hữu.')],[5,11])
    r.h('6.4. Đánh giá giao diện và trải nghiệm')
    r.p('Giao diện sử dụng nhãn tiếng Việt, các thẻ thông tin thống nhất và nút thêm ở vị trí trung tâm. Danh sách hỗ trợ tìm kiếm và lọc để rút ngắn thao tác tra cứu. Biểu đồ giúp nhận biết nhóm chi lớn và mức chi qua các tháng. Ngân sách hiển thị đã chi, còn lại hoặc vượt hạn mức. Chế độ sáng/tối cho phép lựa chọn cách hiển thị phù hợp.')
    r.h('6.5. Hạn chế và phạm vi đánh giá')
    r.p('Kết quả được ghi nhận trên môi trường kiểm thử và máy ảo đã sử dụng. Cần đánh giá bổ sung trên thiết bị thật, màn hình nhỏ, cỡ chữ lớn và với dữ liệu nhiều. Ứng dụng còn nhập liệu thủ công, danh mục cố định, chưa xuất dữ liệu hoặc đồng bộ nhiều thiết bị. Đây là những nội dung được xác định trong hướng phát triển.')

def team(r):
    from build_final_report import MEMBERS
    r.chapter('CHƯƠNG 7. PHÂN CÔNG NHÓM VÀ KẾ HOẠCH BẢO VỆ')
    r.p('Bảng dưới đây là phương án chia phần theo trao đổi của nhóm bốn người. Mỗi người phụ trách cả giao diện, xử lý liên quan, nội dung báo cáo và kiểm thử của phần mình. Thứ tự gán tên theo danh sách bìa chỉ là đề xuất; cần xác nhận lại trước khi coi đây là bảng đóng góp thực tế. Không tự điền phần trăm đóng góp hoặc khẳng định lịch sử thực hiện khi chưa có thông tin xác nhận.')
    tasks=[('Tài khoản và nền tảng database','Đăng ký/đăng nhập/đăng xuất; validator; salt/hash; bốn bảng và quan hệ dữ liệu.','AuthScreen, AuthViewModel, AuthRepository, AccountDao, ExpenseDatabase.','Giải thích Room–SQLite, quan hệ bảng và cách xác minh mật khẩu.'),
        ('Quản lý giao dịch','Biểu mẫu thêm/sửa; chọn loại, danh mục, ngày; kiểm tra tiền; lưu/xóa theo chủ sở hữu.','TransactionEditor, TransactionTile, ExpenseRepository, các truy vấn ghi ExpenseDao.','Demo thêm/sửa/xóa, lỗi nhập và việc giữ dữ liệu khi thao tác thất bại.'),
        ('Tra cứu và tổng quan','Danh sách; tìm kiếm không dấu; lọc tháng và loại; thẻ tổng; điều hướng.','ExpenseScreen, DashboardHeader, BottomNavigation, searchTransactions, summarize.','Giải thích thứ tự lọc, ý nghĩa tổng kỳ và StateFlow cập nhật giao diện.'),
        ('Thống kê và ngân sách','Biểu đồ vòng/cột; công thức tỷ lệ; ngân sách theo tháng và cảnh báo.','ExpenseCharts, BudgetComponents, BudgetStatus, MonthlyBudget, các DAO ngân sách.','Demo các ngưỡng 80%/chạm/vượt và giải thích Canvas.')]
    r.table('Phân công đề xuất cho bốn thành viên',['Thành viên','Phạm vi','Sản phẩm trình bày'],[(n+'\n'+sid,t[0]+'\n'+t[1],t[3]) for (n,sid),t in zip(MEMBERS,tasks)],[3.5,7.5,5])
    for i,((name,sid),t) in enumerate(zip(MEMBERS,tasks),1):
        r.h(f'7.{i}. Thành viên {i}: {name}')
        r.p('Phần phụ trách: '+t[0]+'. '+t[1]); r.p('Tệp/lớp cần đọc: '+t[2]); r.p('Khi trình bày: '+t[3])
        r.p('Nhiệm vụ chung: đối chiếu mô tả trong báo cáo với code, chạy lại kịch bản phần mình, kiểm tra ảnh minh họa và hiểu đường đi của userId. Mỗi người phải biết chức năng của DAO thuộc phần mình; thành viên phụ trách database không làm thay toàn bộ truy vấn cho các phần khác.')
    r.h('7.5. Cân đối khối lượng và phối hợp')
    r.p('Cách chia cân bằng tương đối theo độ khó kiến thức: tài khoản ghép với cấu trúc database; CRUD có nhiều trạng thái biểu mẫu; tra cứu ghép với tổng quan/điều hướng; Canvas ghép với ngân sách. Các lớp dùng chung như ExpenseScreen, ExpenseViewModel và ExpenseDao cần thống nhất phần sửa để tránh xung đột Git. Khi viết báo cáo, từng người bàn giao nội dung và ảnh phần mình thay vì dồn toàn bộ việc tài liệu cho một người.')
    r.h('7.6. Kịch bản bảo vệ khoảng 12–15 phút')
    r.table('Trình tự demo đề xuất',['Thời lượng','Nội dung','Người phụ trách'],[
        ('3 phút','Mục tiêu, đăng nhập và ERD bốn bảng; Room dùng SQLite.','Thành viên 1'),('3–4 phút','Nhập tiền sai để minh họa lỗi; thêm/sửa/xóa một dòng đánh dấu DEMO.','Thành viên 2'),
        ('3 phút','Tìm kiếm không dấu, đổi tháng/loại, giải thích tổng thu/chi và điều hướng.','Thành viên 3'),('3–4 phút','Biểu đồ Canvas, ngân sách và cảnh báo; kết quả kiểm thử và hạn chế.','Thành viên 4')],[2.4,10.4,3.2])
    r.p('Nên tập demo với một tài khoản riêng và dữ liệu có ghi chú DEMO để dễ nhận diện. Không dùng database cá nhân có thông tin nhạy cảm để trình chiếu. Khi nhập một ngày ngoài tháng đang xem, cần chuyển sang tháng đó hoặc Toàn bộ để tìm lại giao dịch; việc chưa thấy trong danh sách hiện tại không có nghĩa thao tác Lưu thất bại.')

def conclusion(r):
    r.chapter('CHƯƠNG 8. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN')
    r.h('8.1. Kết quả đạt được')
    r.p('Đồ án xây dựng ứng dụng quản lý thu/chi cá nhân trên Android với quy trình nhập liệu, lưu trữ, tra cứu và phân tích cơ bản. Dữ liệu được phân tách theo tài khoản cục bộ và lưu bền vững qua Room SQLite. Giao diện Compose có bốn trang và nút thêm trung tâm; các thẻ, biểu mẫu và biểu đồ sử dụng tông màu thống nhất. Ngân sách tháng bổ sung phản hồi cụ thể cho việc kiểm soát chi tiêu.')
    r.p('Về kỹ thuật, nhóm có thể trình bày cách tách View–ViewModel–Repository–DAO, cách Flow đưa dữ liệu lên UI, cấu trúc bốn bảng và công thức vẽ biểu đồ. Bộ kiểm thử cung cấp bằng chứng cho một phần nghiệp vụ và lưu trữ; ảnh thật giúp liên hệ giữa mã nguồn và thao tác sử dụng. Mức độ đáp ứng cuối cùng vẫn cần đối chiếu rubric chính thức của học phần.')
    r.h('8.2. Hướng phát triển')
    for t in ['Bổ sung sao lưu/khôi phục và xuất dữ liệu theo khoảng thời gian.', 'Cho phép người dùng quản lý danh mục, màu và biểu tượng.', 'Tối ưu truy vấn, phân trang và tổng hợp bằng SQL khi dữ liệu tăng.', 'Bổ sung tương tác biểu đồ, hiển thị số tiền khi chọn cột hoặc danh mục.', 'Cải thiện giao diện cho màn hình nhỏ, tablet, cỡ chữ lớn và TalkBack.', 'Nếu mở rộng đồng bộ nhiều thiết bị, thiết kế backend, xác thực và xử lý xung đột riêng.']: r.bullet(t)
    r.chapter('TÀI LIỆU THAM KHẢO')
    refs=[('[1] Android Developers — Save data in a local database using Room','https://developer.android.com/training/data-storage/room'),
        ('[2] Android Developers — Guide to app architecture','https://developer.android.com/topic/architecture'),
        ('[3] Android Developers — ViewModel overview','https://developer.android.com/topic/libraries/architecture/viewmodel'),
        ('[4] Android Developers — Graphics in Compose','https://developer.android.com/develop/ui/compose/graphics/draw/overview'),
        ('[5] Repository mã nguồn QLCT_LTDD','https://github.com/nhuvu988/QLCT_LTDD')]
    for title,url in refs: r.p(title+'\n'+url+'\nNgày đối chiếu: 21/10/2026.')
    r.p('Các tài liệu chính thức được dùng để đối chiếu khái niệm Room, kiến trúc, ViewModel và Canvas. Phiên bản thư viện, thuật toán, bảng dữ liệu và kết quả kiểm thử trong báo cáo lấy từ mã nguồn và bằng chứng chạy của repository, không suy ra từ phiên bản ví dụ trên trang tài liệu.')

def write_all(r):
    analysis(r); design(r); implementation(r); interface(r); testing(r); team(r); conclusion(r)
