"""Generate the budget/search report supplement (python-docx and matplotlib)."""
from pathlib import Path
from docx import Document
from docx.shared import Inches
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/diagrams/budget_search'
OUT.mkdir(parents=True,exist_ok=True)

def diagram(name,title,boxes,edges):
    fig,ax=plt.subplots(figsize=(10,6))
    ax.set(xlim=(0,10),ylim=(0,6)); ax.axis('off'); ax.set_title(title,fontsize=16,pad=18)
    for key,(x,y,w,h,text) in boxes.items():
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.08',facecolor='#eee9fa',edgecolor='#7962c8'))
        ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=10,wrap=True)
    for a,b,label in edges:
        x,y,w,h,_=boxes[a]; xx,yy,ww,hh,_=boxes[b]
        start=(x+w/2,y) if yy<y else (x+w,y+h/2)
        end=(xx+ww/2,yy+hh) if yy<y else (xx,yy+hh/2)
        ax.annotate('',xy=end,xytext=start,arrowprops=dict(arrowstyle='->',color='#534875',lw=1.5))
        if label: ax.text((start[0]+end[0])/2,(start[1]+end[1])/2,label,fontsize=9,ha='center',bbox=dict(facecolor='white',edgecolor='none',pad=1))
    fig.savefig(OUT/(name+'.png'),dpi=160,bbox_inches='tight')
    fig.savefig(OUT/(name+'.svg'),bbox_inches='tight'); plt.close(fig)
    svg = OUT/(name+'.svg')
    svg.write_text('\n'.join(line.rstrip() for line in svg.read_text(encoding='utf-8').splitlines())+'\n',encoding='utf-8')

diagram('use_case','Use Case — Ngân sách và tìm kiếm',{
 'user':(0.2,2.6,2,1,'Người dùng\nđã đăng nhập'),
 'budget':(4,4.1,5.3,1,'Thiết lập / sửa / xóa ngân sách tháng'),
 'status':(4,2.6,5.3,1,'Xem còn lại và cảnh báo ngân sách'),
 'search':(4,1.1,5.3,1,'Tìm kiếm giao dịch\nkết hợp bộ lọc tháng, Thu/Chi')},
 [('user','budget',''),('user','status',''),('user','search','')])
diagram('flow','Luồng thiết lập ngân sách và tìm kiếm',{
 'a':(.3,4.5,4.2,.75,'Chọn tháng → Nhập hạn mức'),
 'b':(.3,2.9,4.2,.9,'Kiểm tra số nguyên VND > 0\nSai: giữ biểu mẫu và báo lỗi'),
 'c':(.3,1.3,4.2,.9,'Lưu (user_id, month) vào Room\nFlow cập nhật tiến trình / cảnh báo'),
 'd':(5.4,4.5,4.2,.75,'Tab Giao dịch → Nhập từ khóa'),
 'e':(5.4,2.9,4.2,.9,'Lọc giao dịch của tài khoản\ntheo tháng và loại Thu/Chi'),
 'f':(5.4,1.3,4.2,.9,'Chuẩn hóa dấu và chữ hoa/thường\nKhớp ghi chú, danh mục → kết quả')},
 [('a','b',''),('b','c','Hợp lệ'),('d','e',''),('e','f','')])
diagram('erd','ERD bổ sung — ngân sách thuộc tài khoản',{
 'a':(.3,3.7,3.1,1.3,'local_account\nPK: id\nusername, salt, passwordHash'),
 'b':(5.6,3.7,3.9,1.3,'budgets\nPK: (user_id, month)\nFK: user_id → local_account.id\namount: số nguyên VND'),
 'c':(5.6,.8,3.9,1.4,'transactions\nPK: id\nFK: user_id → local_account.id\namount, date, note, category_id')},
 [('a','b','1 : N'),('a','c','1 : N')])
diagram('class','Class Diagram — các thành phần bổ sung',{
 'a':(.3,4.3,3.8,1,'ExpenseScreen\nBudgetEditor / MonthlyBudgetCard\nÔ tìm kiếm giao dịch'),
 'b':(.3,2.4,3.8,1,'ExpenseViewModel\nsaveBudget(), deleteBudget()\nsetQuery(), state: StateFlow'),
 'c':(5.8,2.4,3.8,1,'ExpenseRepository\nbudgets: Flow\nsaveBudget(), deleteBudget()'),
 'd':(5.8,.4,3.8,1,'ExpenseDao / MonthlyBudget\nobserveBudgets(), @Upsert\n(userId, month, amount)'),
 'e':(.3,.4,3.8,1,'BudgetStatus / searchTransactions\n80%, 100%, vượt mức\nChuẩn hóa dấu, lọc từ khóa')},
 [('a','b',''),('b','c',''),('b','e',''),('c','d','')])

doc=Document()
doc.add_heading('PHỤ LỤC — NGÂN SÁCH THÁNG VÀ TÌM KIẾM GIAO DỊCH',0)
doc.add_paragraph('Đồ án QLCT · Cập nhật 07/10/2026 · Bổ sung cho báo cáo Compose–MVVM–Room hiện có.')
doc.add_heading('1. Phạm vi và quy tắc',level=1)
for text in [
 'Ngân sách là hạn mức tổng chi của một tháng, riêng từng tài khoản. Phiên bản này chưa có hạn mức theo danh mục hoặc thông báo nền.',
 'Một tài khoản chỉ có một ngân sách cho mỗi tháng. Số tiền từ 1 đến 999.999.999.999 VND. Lưu lại cùng tháng cập nhật hạn mức; tháng chưa đặt ngân sách hiển thị trạng thái chưa thiết lập.',
 'Tổng chi lấy toàn bộ khoản EXPENSE trong tháng đang chọn của tài khoản đăng nhập. Không phụ thuộc từ khóa tìm kiếm, bộ lọc Thu/Chi hoặc tùy chọn xem toàn bộ thời gian.',
 'Cảnh báo: dưới 80% bình thường; từ 80% đến dưới 100% màu vàng; đúng 100% đã dùng hết; trên 100% báo vượt mức. Xóa ngân sách có xác nhận và không xóa giao dịch.',
 'Tìm kiếm ghi chú và tên danh mục trong tab Giao dịch, không phân biệt hoa/thường hoặc dấu tiếng Việt. Mọi từ trong truy vấn phải xuất hiện trong ghi chú/danh mục. Từ khóa rỗng khôi phục danh sách theo bộ lọc hiện hành.',
 'Kết quả tìm kiếm tôn trọng tháng và loại giao dịch; không thay đổi tổng thu/chi, biểu đồ hay ngân sách. SavedStateHandle giữ từ khóa khi tạo lại màn hình.'
]: doc.add_paragraph(text)
doc.add_picture(str(OUT/'use_case.png'),width=Inches(6.2))
doc.add_heading('2. Luồng thao tác',level=1)
doc.add_paragraph('Tổng quan/Thống kê → chọn tháng → Đặt/Sửa ngân sách → nhập số tiền → Lưu. Khi sửa/xóa/thêm hoặc chuyển ngày giao dịch, Room Flow tính lại tổng chi và trạng thái ngân sách.')
doc.add_paragraph('Giao dịch → nhập từ khóa, ví dụ “an trua” → xem số kết quả → kết hợp bộ lọc tháng/Thu/Chi → chạm kết quả để sửa hoặc xóa. Khi không khớp, hiển thị gợi ý đổi từ khóa hoặc bộ lọc.')
doc.add_picture(str(OUT/'flow.png'),width=Inches(6.2))
doc.add_heading('3. Database và kiến trúc',level=1)
doc.add_paragraph('Room phiên bản 6 thêm bảng budgets. Migration 5→6 chỉ tạo bảng mới, giữ tài khoản, giao dịch, danh mục và mật khẩu. Khóa chính ghép (user_id, month) ngăn trùng ngân sách, khóa ngoại user_id liên kết local_account.id. Không dùng fallbackToDestructiveMigration.')
doc.add_picture(str(OUT/'erd.png'),width=Inches(6.2))
doc.add_picture(str(OUT/'class.png'),width=Inches(6.2))
doc.add_heading('4. Kiểm thử và giới hạn',level=1)
doc.add_paragraph('Nguồn kiểm thử: BudgetAndSearchTest.kt, MonthlyBudgetTest.kt và docs/tests/test_budget_migration.py. Kết quả chạy chính xác được lưu trong docs/budget_search_validation.txt. Sơ đồ trong phụ lục là mô hình thiết kế, không phải ảnh chụp màn hình chạy thật.')
table=doc.add_table(rows=1, cols=2); table.style='Table Grid'
table.rows[0].cells[0].text='Tình huống'; table.rows[0].cells[1].text='Kết quả mong đợi'
for a,b in [
 ('Hạn mức 1.000, chi 799 / 800 / 1.000 / 1.200','Bình thường / cảnh báo / hết / vượt 200'),
 ('Số tiền 0, âm, vượt giới hạn hoặc tháng sai','Từ chối và giữ dữ liệu nhập'),
 ('Hai tài khoản cùng tháng','Ngân sách tách riêng'),
 ('Đóng/mở database; cập nhật app từ v5','Dữ liệu vẫn còn, không tạo trùng'),
 ('Tìm “an truong”, “ĐỔ XĂNG”, khoảng trắng','Khớp không dấu, hoa/thường; khoảng trắng trả danh sách'),
 ('Sửa/xóa giao dịch hoặc chuyển sang tháng khác','Tổng chi và ngân sách tự cập nhật'),
 ('Tìm kiếm khi lọc Thu/Chi hoặc toàn bộ thời gian','Chỉ lọc danh sách, không giảm tổng chi ngân sách'),
 ('Xóa hạn mức','Giao dịch và ngân sách tháng/tài khoản khác giữ nguyên'),
 ('Thử trên thiết bị','Cần kiểm tra bàn phím, xoay màn hình và tiến trình thực tế')
]:
    cells=table.add_row().cells; cells[0].text=a; cells[1].text=b
doc.save(ROOT/'docs/Phu_luc_Ngan_sach_Tim_kiem.docx')
print('Generated Word supplement and four diagrams (PNG + SVG)')
