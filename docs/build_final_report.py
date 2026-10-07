"""Build the detailed report from current source, evidence, and editable diagrams.

Dependencies are kept in .gradle/report-libs; this script never changes app data.
"""
from pathlib import Path
import sys, json, re
from zipfile import ZipFile
from xml.etree import ElementTree as ET
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.gradle/report-libs'))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

OUT = ROOT / 'docs/Bao_cao_Do_an_QLCT_Compose_MVVM_Room.docx'
MEMBERS = [('Nguyễn Như Vũ','25610008'),('Nguyễn Văn Vũ','25810053'),('Nguyễn Thiên Vọng','25810050'),('Nguyễn Văn Phước','25810035')]

def field(paragraph, instruction, placeholder=''):
    run=paragraph.add_run()
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    code=OxmlElement('w:instrText'); code.set(qn('xml:space'),'preserve'); code.text=instruction
    separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
    text=OxmlElement('w:t'); text.text=placeholder
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    for item in (begin,code,separate,text,end): run._r.append(item)

class Report:
    def __init__(self):
        self.doc=Document(); self.figures=[]; self.tables=[]; self.bookmark_id=1
        sec=self.doc.sections[0]
        sec.page_width=Cm(21); sec.page_height=Cm(29.7)
        sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(3); sec.right_margin=Cm(2)
        sec.header_distance=Cm(.9); sec.footer_distance=Cm(.9)
        sec.different_first_page_header_footer=True
        for name in ('Normal','Body Text','List Bullet','List Number'):
            s=self.doc.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(13)
            s._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
            s.paragraph_format.space_after=Pt(7); s.paragraph_format.line_spacing=1.3
        for name,size in [('Title',23),('Heading 1',16),('Heading 2',14),('Heading 3',13)]:
            s=self.doc.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(size); s.font.bold=True
            s.font.color.rgb=RGBColor.from_string('504079')
            s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(8)
            s.paragraph_format.keep_with_next=True
        self.doc.styles['Caption'].font.name='Times New Roman'; self.doc.styles['Caption'].font.size=Pt(11)
        self.doc.styles['Caption'].font.color.rgb=RGBColor.from_string('555555')
        self.doc.styles['Caption'].paragraph_format.space_after=Pt(8)
        head=sec.header.paragraphs[0]; head.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        rr=head.add_run('QLCT  |  Đồ án Lập trình di động'); rr.font.size=Pt(9); rr.font.color.rgb=RGBColor.from_string('777777')
        foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field(foot,' PAGE ','1')
        settings=self.doc.settings.element
        update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
        self.doc.core_properties.title='Báo cáo đồ án: Ứng dụng quản lý chi tiêu cá nhân — QLCT'
        self.doc.core_properties.author='; '.join(x[0] for x in MEMBERS)
        self.doc.core_properties.subject='Jetpack Compose • MVVM • Room • Android'
        self.doc.core_properties.keywords='QLCT, Compose, MVVM, Room, báo cáo đồ án'

    def p(self,text='',bold=False):
        p=self.doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.widow_control=True
        p.add_run(text).bold=bold
        return p

    def h(self,text,level=2):
        if text == 'MỤC LỤC': return self.doc.add_paragraph(text,style='Title')
        return self.doc.add_heading(text,level)

    def chapter(self,text):
        self.doc.add_page_break(); self.h(text,1)

    def bullet(self,text): return self.doc.add_paragraph(text,style='List Bullet')

    def steps(self,items):
        for i,item in enumerate(items,1): self.p(f'{i}. {item}')

    def code(self,text):
        p=self.doc.add_paragraph()
        p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(10)
        p.paragraph_format.line_spacing=1.05
        r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(9)
        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F3F1F7'); p._p.get_or_add_pPr().append(shade)
        return p

    def bookmark(self,p,name):
        ident=str(self.bookmark_id); self.bookmark_id+=1
        begin=OxmlElement('w:bookmarkStart'); begin.set(qn('w:id'),ident); begin.set(qn('w:name'),name)
        end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),ident)
        p._p.insert(0,begin); p._p.append(end)

    def table(self,caption,headers,rows,widths=None):
        num=len(self.tables)+1; label=f'Bảng {num}. {caption}'
        cp=self.doc.add_paragraph(label,style='Caption'); cp.paragraph_format.keep_with_next=True
        self.bookmark(cp,f'table_{num}'); self.tables.append((label,f'table_{num}'))
        t=self.doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False
        if widths is None: widths=[16/len(headers)]*len(headers)
        for c,w in zip(t.columns,widths): c.width=Cm(w)
        for c,text,w in zip(t.rows[0].cells,headers,widths):
            c.width=Cm(w); c.text=str(text)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8E2F3'); c._tc.get_or_add_tcPr().append(sh)
            for r in c.paragraphs[0].runs: r.bold=True
        repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for c,value,w in zip(cells,row,widths): c.width=Cm(w); c.text=str(value)
        for row in t.rows:
            cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3); p.paragraph_format.line_spacing=1.12
                    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(10.5)
        self.doc.add_paragraph().paragraph_format.space_after=Pt(1)
        return t

    def image(self,path,caption,width=15.5):
        path=ROOT/path
        if not path.exists(): raise FileNotFoundError(path)
        with Image.open(path) as im: iw,ih=im.size
        width=min(width,20*iw/ih)
        p=self.doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(path),width=Cm(width))
        num=len(self.figures)+1; label=f'Hình {num}. {caption}'
        cp=self.doc.add_paragraph(label,style='Caption'); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.bookmark(cp,f'figure_{num}'); self.figures.append((label,f'figure_{num}'))

    def pair_images(self,left,left_caption,right,right_caption):
        t=self.doc.add_table(rows=1,cols=2); t.autofit=False
        for cell,path,caption in zip(t.rows[0].cells,[left,right],[left_caption,right_caption]):
            cell.width=Cm(8)
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT/path),width=Cm(7))
            num=len(self.figures)+1; label=f'Hình {num}. {caption}'
            cp=cell.add_paragraph(label,style='Caption'); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            self.bookmark(cp,f'figure_{num}'); self.figures.append((label,f'figure_{num}'))
        cant=OxmlElement('w:cantSplit'); t.rows[0]._tr.get_or_add_trPr().append(cant)

    def front(self):
        p=self.doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run('[TÊN TRƯỜNG]\n[KHOA / BỘ MÔN]'); r.bold=True; r.font.size=Pt(15)
        self.doc.add_paragraph('\n')
        p=self.doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run('BÁO CÁO ĐỒ ÁN\nMÔN LẬP TRÌNH DI ĐỘNG'); r.bold=True; r.font.size=Pt(21)
        p=self.doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run('ỨNG DỤNG QUẢN LÝ\nCHI TIÊU CÁ NHÂN'); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string('67509F')
        p=self.doc.add_paragraph('QLCT — Android | Jetpack Compose | MVVM | Room'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.doc.add_paragraph()
        self.p('Giảng viên hướng dẫn: [Bổ sung họ tên giảng viên]')
        self.p('Lớp: [Bổ sung tên lớp]     •     Nhóm: 4 thành viên')
        self.table('Danh sách thành viên nhóm',['STT','Họ và tên','Mã sinh viên'],[(str(i),n,s) for i,(n,s) in enumerate(MEMBERS,1)],[1.2,9.8,5])
        p=self.doc.add_paragraph('Năm học: [Bổ sung]\nBản báo cáo theo mã nguồn ngày 06/10/2026'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.doc.add_page_break(); self.h('THÔNG TIN BÁO CÁO',1)
        self.p('Báo cáo trình bày ứng dụng QLCT đã triển khai trong repository QLCT_LTDD. Các mô tả nghiệp vụ, lớp, bảng dữ liệu và hình ảnh được đối chiếu với phiên bản Jetpack Compose, MVVM và Room version 2. Báo cáo Basic và đề cương cũ trong docs/ là tài liệu lịch sử, không phải mô tả kiến trúc hiện tại.')
        self.p('Các mục tên trường, khoa, lớp, giảng viên và năm học trên trang bìa chưa được cung cấp tại thời điểm biên soạn nên được để dưới dạng chỗ trống có nhãn. Nhóm cần hoàn thiện các mục này trước khi nộp. Phân công trong báo cáo là phương án đề xuất để tổ chức thuyết trình, không khẳng định người nào đã thực hiện một phần mã cụ thể.')
        self.h('Tóm tắt đồ án',2)
        self.p('Ứng dụng giúp người dùng ghi nhận khoản thu và khoản chi, xem lịch sử, chọn thời gian, lọc loại giao dịch và quan sát tổng quan tài chính. Dữ liệu được lưu cục bộ bằng Room trên SQLite. Giao diện sử dụng Compose Material 3 với ba trang Tổng quan, Giao dịch và Thống kê; biểu đồ vòng thể hiện phân bổ khoản chi theo danh mục, biểu đồ cột thể hiện tổng chi trong sáu tháng. Kiến trúc MVVM tách giao diện, trạng thái và truy cập dữ liệu để dễ kiểm thử và bảo trì.')
        self.p('Ứng dụng không cần tài khoản hoặc máy chủ để thực hiện các chức năng hiện có. Số tiền lưu dạng Long, luôn lớn hơn 0; loại INCOME hoặc EXPENSE quyết định ý nghĩa thu/chi. Có migration bảo toàn danh mục và giao dịch từ database cũ version 1. Các bằng chứng kiểm thử gồm unit test, instrumentation test, kiểm tra SQL migration trên máy tính và một số luồng thao tác trực tiếp trên máy ảo.')
        self.h('Từ khóa',2); self.p('Android; Kotlin; Jetpack Compose; MVVM; Room; SQLite; quản lý chi tiêu; kiểm thử; migration.')
        self.doc.add_page_break(); self.h('MỤC LỤC',1)
        field(self.doc.add_paragraph(),' TOC \\o "1-3" \\h \\z \\u ','Mục lục tự động — cập nhật bằng Ctrl+A, F9 trong Word.')
        self.doc.add_page_break(); self.h('DANH MỤC HÌNH',1); self.figure_index=self.doc.add_paragraph()
        self.doc.add_page_break(); self.h('DANH MỤC BẢNG',1); self.table_index=self.doc.add_paragraph()
        self.doc.add_page_break(); self.h('DANH MỤC TỪ VIẾT TẮT',1)
        self.table('Từ viết tắt và ý nghĩa',['Thuật ngữ','Giải thích'],[
            ('UI','User Interface — giao diện người dùng'),('MVVM','Model – View – ViewModel'),('DAO','Data Access Object — lớp/interface truy cập dữ liệu'),('CRUD','Create, Read, Update, Delete — thêm, đọc, sửa, xóa'),('ERD','Entity Relationship Diagram — sơ đồ quan hệ thực thể'),('PK / FK','Primary Key / Foreign Key — khóa chính / khóa ngoại'),('VND','Đồng Việt Nam; sử dụng số nguyên trong ứng dụng'),('Flow / StateFlow','Luồng dữ liệu bất đồng bộ / luồng giữ trạng thái mới nhất'),('KSP','Kotlin Symbol Processing — sinh mã Room'),('AVD','Android Virtual Device — cấu hình máy ảo Android'),('SDK / API','Software Development Kit / Application Programming Interface'),('IME','Input Method Editor — phương thức nhập/bàn phím Android')],[3.2,12.8])

    def finish(self):
        for anchor,items in [(self.figure_index,self.figures),(self.table_index,self.tables)]:
            for label,bookmark in items:
                p=self.doc.add_paragraph(); p.paragraph_format.line_spacing=1.1; p.paragraph_format.space_after=Pt(5)
                p.add_run(label+'  '); field(p,f' PAGEREF {bookmark} \\h ','—')
                anchor._p.addprevious(p._p)
        self.doc.save(OUT)
        with ZipFile(OUT) as z:
            assert z.testzip() is None
            for name in z.namelist():
                if name.endswith(('.xml','.rels')): ET.fromstring(z.read(name))
            text=z.read('word/document.xml').decode('utf-8')
            for _,msv in MEMBERS: assert msv in text
            for name in z.namelist():
                if name.startswith('word/media/'):
                    assert len(z.read(name))>1000
        audit={'file':str(OUT),'paragraphs':len(self.doc.paragraphs),'tables':len(self.tables),'figures':len(self.figures),'members':MEMBERS,'word_count':sum(len(p.text.split()) for p in self.doc.paragraphs)+sum(len(c.text.split()) for t in self.doc.tables for row in t.rows for c in row.cells),'validated':'DOCX ZIP, XML, embedded images, student IDs'}
        (ROOT/'docs/report_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps(audit,ensure_ascii=True))

def main():
    from report_diagrams import generate
    generate()
    from report_sections_analysis import write_analysis
    from report_sections_design import write_design
    from report_sections_testing import write_testing
    report=Report(); report.front()
    write_analysis(report); write_design(report); write_testing(report)
    report.finish()

if __name__ == '__main__': main()
